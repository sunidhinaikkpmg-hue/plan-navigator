"""
connect.py – loads Plan Navigator Data.xlsx once at startup and exposes
sheet-level accessors used by the API routers.
"""
from __future__ import annotations

import json
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl

_XLSX_PATH = Path(__file__).parent / "Plan Navigator Data.xlsx"


@lru_cache(maxsize=1)
def _workbook() -> openpyxl.Workbook:
    return openpyxl.load_workbook(_XLSX_PATH, read_only=True, data_only=True)


def _sheet_rows(sheet_name: str) -> list[dict[str, Any]]:
    """Return all data rows from *sheet_name* as a list of dicts."""
    wb = _workbook()
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = rows[0]
    result = []
    for row in rows[1:]:
        record: dict[str, Any] = {}
        for header, value in zip(headers, row):
            if header is None:
                continue
            # Normalise datetime objects to ISO strings
            if hasattr(value, "isoformat"):
                value = value.isoformat()
            record[header] = value
        # Skip completely blank rows (all None / empty after header drop)
        if any(v is not None for v in record.values()):
            result.append(record)
    return result


def _sheet_headers(sheet_name: str) -> list[str]:
    """Return header names for *sheet_name* preserving workbook order."""
    wb = _workbook()
    ws = wb[sheet_name]
    first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers: list[str] = []
    for header in first_row:
        if header is None:
            continue
        header_text = str(header).strip()
        if header_text:
            headers.append(header_text)
    return headers


def _is_table_overview_sheet(sheet_name: str) -> bool:
    normalized = "".join(ch for ch in sheet_name.lower() if ch.isalnum())
    return normalized == "tableoverview"


def get_plan_sheet_data(
    sheet_name: str,
    plan_id: str,
    *,
    plan_id_column: str = "plan_id",
    fields: list[str] | tuple[str, ...] | None = None,
) -> list[dict[str, Any]]:
    """
    Return rows for a plan from any sheet.

    - If ``fields`` is provided, only those columns are extracted when present.
    - If ``fields`` is None, full row payload is returned.
    """
    rows = [r for r in _sheet_rows(sheet_name) if r.get(plan_id_column) == plan_id]
    if fields is None:
        return rows

    extracted: list[dict[str, Any]] = []
    for row in rows:
        extracted.append({field: row.get(field) for field in fields if field in row})
    return extracted


def get_data_schema(offset: int = 0, limit: int = 10) -> dict[str, Any]:
    """Return workbook sheets with columns and paginated row payloads."""
    safe_offset = max(0, offset)
    safe_limit = max(1, limit)

    wb = _workbook()

    sheets: list[dict[str, Any]] = []
    total_rows = 0
    returned_rows = 0
    for sheet_name in wb.sheetnames:
        if _is_table_overview_sheet(sheet_name):
            continue

        rows = _sheet_rows(sheet_name)
        paginated_rows = rows[safe_offset : safe_offset + safe_limit]
        sheets.append(
            {
                "sheet_name": sheet_name,
                "columns": _sheet_headers(sheet_name),
                "row_count": len(rows),
                "returned_row_count": len(paginated_rows),
                "rows": paginated_rows,
            }
        )
        total_rows += len(rows)
        returned_rows += len(paginated_rows)

    return {
        "offset": safe_offset,
        "limit": safe_limit,
        "total_sheets": len(sheets),
        "total_rows": total_rows,
        "returned_rows": returned_rows,
        "sheets": sheets,
    }


# ---------------------------------------------------------------------------
# Public helpers
# ---------------------------------------------------------------------------

def get_test_results(plan_id: str) -> list[dict]:
    return [r for r in _sheet_rows("test_results") if r.get("plan_id") == plan_id]


def get_diagnostic_tests() -> dict[str, dict]:
    """Keyed by diagnostic_tests_id for O(1) join."""
    return {r["diagnostic_tests_id"]: r for r in _sheet_rows("diagnostic_tests")}


def get_peer_benchmarks() -> dict[str, dict]:
    """Keyed by test_id for O(1) join with diagnostic/test result rows."""
    return {r["test_id"]: r for r in _sheet_rows("peer_benchmarks") if r.get("test_id")}


def get_plan(plan_id: str) -> dict[str, Any] | None:
    """Return a single plan row from the plans sheet."""
    for row in _sheet_rows("plans"):
        if row.get("plan_id") == plan_id:
            return row
    return None


def get_recommendations(plan_id: str) -> list[dict]:
    return [r for r in _sheet_rows("recommendations") if r.get("plan_id") == plan_id]


def get_campaigns(plan_id: str) -> list[dict]:
    results = []
    for r in _sheet_rows("education_campaigns"):
        raw = r.get("plan_id", "")
        try:
            ids = json.loads(raw) if isinstance(raw, str) and raw.startswith("[") else [raw]
        except (ValueError, TypeError):
            ids = [raw]
        if plan_id in ids:
            r = dict(r)
            r["plan_id"] = ids  # normalise to list
            results.append(r)
    return results


def get_documents(plan_id: str) -> list[dict]:
    rows = []
    for r in _sheet_rows("document"):
        # The document sheet has a duplicate header row — skip it
        if r.get("id") == "id":
            continue
        if r.get("plan_id") == plan_id:
            rows.append(r)
    return rows


def get_usage_events() -> list[dict]:
    return _sheet_rows("usage_events")


def get_ai_interactions() -> list[dict]:
    return _sheet_rows("ai_interactions")


def get_payroll_files(plan_id: str) -> list[dict]:
    return [r for r in _sheet_rows("payroll_files") if r.get("plan_id") == plan_id]


def get_compliance_filings(plan_id: str) -> list[dict]:
    return [r for r in _sheet_rows("compliance_filings") if r.get("plan_id") == plan_id]
